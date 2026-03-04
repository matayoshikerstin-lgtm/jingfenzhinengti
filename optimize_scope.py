import docx
from docx.shared import Pt

def optimize_scope():
    file_path = r'test_plan\经分智能体测试方案2.0.docx'
    try:
        doc = docx.Document(file_path)
    except Exception as e:
        print(f"Error opening document: {e}")
        return

    # Find the section "二. 测试范围"
    start_idx = -1
    for i, p in enumerate(doc.paragraphs):
        if '二. 测试范围' in p.text or '二、' in p.text:
            start_idx = i
            break
            
    if start_idx == -1:
        print("未找到 二. 测试范围")
        return

    # Delete the next 3 paragraphs (输入, 方式, 问题分类)
    # Be careful not to delete tables or other sections
    paragraphs_to_remove = []
    for i in range(start_idx + 1, start_idx + 4):
        if i < len(doc.paragraphs) and '三.' not in doc.paragraphs[i].text:
            paragraphs_to_remove.append(doc.paragraphs[i])

    target_p = doc.paragraphs[start_idx]
    target_p.text = '二. 测试范围与边界定界 (Test Scope & Boundaries)'
    
    p1 = target_p.insert_paragraph_before('1. 评测输入源 (Input Vectors)：')
    p1.runs[0].bold = True
    target_p.insert_paragraph_before('依托真实的业务系统快照构建动态测试集（测试数据.xlsx），提取出包含 80 道高拟真度复合型经营问题的全量题库，实现对基础事实提取与高阶商业逻辑推演的双重压测。')

    p2 = target_p.insert_paragraph_before('2. 触发与交互方式 (Execution Modality)：')
    p2.runs[0].bold = True
    target_p.insert_paragraph_before('摒弃传统的人工点选，全面采用 RPA (Robotic Process Automation) 引擎直接挂载智能体的前端 UI。通过自动注入问题集并实时监听抓取流式响应结果，实现闭环全异步跑批。')

    p3 = target_p.insert_paragraph_before('3. 评测维度分类 (Taxonomy of Queries)：')
    p3.runs[0].bold = True
    target_p.insert_paragraph_before('为全面摸底大模型的能力边界，将测试题库解耦为“确定性事实”与“主观性推演”两大阵列：')

    # Remove the old paragraphs
    for p in paragraphs_to_remove:
        p._element.getparent().remove(p._element)

    # Now update Table 0 headers and content slightly to match the professional tone
    table = doc.tables[0]
    
    table.rows[0].cells[0].text = '问题阵列'
    table.rows[0].cells[1].text = '能力定义与指令范式'
    table.rows[0].cells[2].text = '题库配比'
    table.rows[0].cells[3].text = '核心锚点'
    
    table.rows[1].cells[0].text = '非开放问题\n(Fact-Based)'
    table.rows[1].cells[1].text = '客观事实检索，具备绝对唯一的 Ground Truth。\n示例：“净利率最高的合同是哪一个？”'
    table.rows[1].cells[2].text = '50 道\n(高阶逻辑)'
    table.rows[1].cells[3].text = '数据提取精准度、跨表聚合能力'

    table.rows[2].cells[0].text = '开放问题\n(Insight-Based)'
    table.rows[2].cells[1].text = '主观归纳与经营诊断，需提供三段式推理引擎输出。\n示例：“为什么杨帆利润这么低？请给出管理对策。”'
    table.rows[2].cells[2].text = '30 道\n(灵魂拷问)'
    table.rows[2].cells[3].text = '事实零幻觉、商业逻辑严密性、对策可行性'

    try:
        doc.save(file_path)
        print("【二. 测试范围】已完成极致优化！")
    except Exception as e:
        print(f"保存文件失败: {e}")

if __name__ == '__main__':
    optimize_scope()