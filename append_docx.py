import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

def append_llm_judge():
    file_path = r'test_plan\经分智能体测试方案2.0.docx'
    try:
        doc = docx.Document(file_path)
    except Exception as e:
        print(f"Error opening document: {e}")
        return

    # 添加新段落和标题
    doc.add_page_break()
    p_h1 = doc.add_paragraph()
    r_h1 = p_h1.add_run('附录：开放问题评估与打分标准（LLM-as-a-Judge 自动化打分法）')
    r_h1.bold = True
    r_h1.font.size = Pt(16)
    
    # 添加背景说明
    p1 = doc.add_paragraph('鉴于大模型（智能体）的输出为自由文本，且通常不会严格按照设定的“三段式”固定格式输出。为保证评分的客观性、一致性与高效性，测试方案引入 ')
    run_bold = p1.add_run('LLM-as-a-Judge（用大模型当裁判）')
    run_bold.bold = True
    p1.add_run(' 的自动化评分机制。')
    
    p_h2_1 = doc.add_paragraph()
    r_h2_1 = p_h2_1.add_run('1. 核心评估理念')
    r_h2_1.bold = True
    r_h2_1.font.size = Pt(14)
    
    doc.add_paragraph('评分时忽略智能体回答的排版格式（如是否带小标题），纯看其生成的语义是否涵盖了基准答案中的三大核心考核点：\n'
                      '• 事实与数据层：是否找对底层数据，有无幻觉。\n'
                      '• 归纳与分析层：是否把数据转化为业务洞察。\n'
                      '• 落地与建议层：是否给出可执行的管理建议。')

    p_h2_2 = doc.add_paragraph()
    r_h2_2 = p_h2_2.add_run('2. LLM 裁判指令 (Prompt) 模板')
    r_h2_2.bold = True
    r_h2_2.font.size = Pt(14)
    prompt_text = """【系统设定】
你是一个资深的企业经营分析总监，现在你需要给实习生（经分智能体）生成的业务分析报告打分。

【输入信息】
1. 用户提问：[测试用例中的开放问题]
2. 参考基准答案：[测试方案中提供的带数据和建议的三段式标准答案]
3. 智能体实际回答：[被测经分智能体生成的实际长文本]

【打分规则（满分 5 分）】
请忽略智能体回答的排版格式，纯看语义是否涵盖了基准答案的核心考点：
- 准确性（2分）：实际回答中引用的业务数据、数值和客户名称，是否与基准答案一致，有无编造（幻觉）？
- 洞察力（1.5分）：实际回答是否指出了基准答案中提到的业务痛点或现象根因？
- 建议有效性（1.5分）：实际回答是否给出了与基准答案方向一致或同样具有业务指导意义的管理建议？

【输出格式】
1. 得分：[X / 5分]
2. 扣分理由：（如果满分则填“涵盖全面，数据精准”）"""
    
    p_prompt = doc.add_paragraph()
    run_prompt = p_prompt.add_run(prompt_text)
    run_prompt.font.name = 'Courier New'
    run_prompt.font.size = Pt(10)

    p_h2_3 = doc.add_paragraph()
    r_h2_3 = p_h2_3.add_run('3. 自动化测试实施路径')
    r_h2_3.bold = True
    r_h2_3.font.size = Pt(14)
    
    doc.add_paragraph('步骤一：通过 RPA 脚本或 API 批量拉取智能体对 80 道开放题的实际回答，并落盘至 result.xlsx 的“智能体回复”列。')
    doc.add_paragraph('步骤二：编写评测脚本，循环读取“问题 + 标准答案 + 智能体回复”，将其拼接到上述 Prompt 模板中。')
    doc.add_paragraph('步骤三：调用更高阶的大模型（如 GPT-4o 或 Claude 3.5 Sonnet）接口，让其以裁判身份输出分数与扣分理由。')
    doc.add_paragraph('步骤四：汇总分数，计算整体平均分以及“准确性、洞察力、建议”等细分维度的达标率。')

    try:
        doc.save(file_path)
        print(f"成功将 LLM-as-a-Judge 自动化打分方案追加到 {file_path}")
    except Exception as e:
        print(f"保存文件失败: {e} (请确认文件未在 Word 中打开)")

if __name__ == '__main__':
    append_llm_judge()
